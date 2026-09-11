from __future__ import annotations

import math
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# =============================================================================
# USER PATHS
# =============================================================================
BASE = Path(r"D:\Headway")

T5_PATH = BASE / "Tables" / "T5_41_Sequential_AFT_Selection_VIF.xlsx"
T6_PATH = BASE / "Tables" / "T6_Final_AFT_Bootstrap_CoxSnell.xlsx"
DATA_PATH = BASE / "data3.xlsx"

OUTPUT_XLSX = BASE / "final run" / "tables" / "Table8_Final_AFT_Models.xlsx"
OUTPUT_CSV = BASE / "final run" / "tables" / "Table8_Final_AFT_Models.csv"

PAIR_ORDER = [
    "BTW_following_4W",
    "BTW_following_MT_3W",
    "BTW_following_NMT_3W",
    "BTW_following_MT_2W",
    "BTW_following_NMT_2W",
    "PR_following_4W",
    "PR_following_MT_3W",
    "PR_following_NMT_3W",
]

COVARIATE_LABELS = {
    "Speed_Difference": "Speed difference",
    "Target_Speed_km/hr": "Target speed",
    "Leading_Speed_km/hr": "Leader speed",
    "Off_centeredness": "Off-centeredness",
    "Occupancy": "Occupancy",
    "Flow_pcu/hr/m": "Traffic flow",
    "Site": "Site",
}

# =============================================================================
# HELPERS
# =============================================================================
def is_blank(value: Any) -> bool:
    return pd.isna(value) or str(value).strip() in {"", "nan", "None"}


def as_float(value: Any, label: str) -> float:
    if is_blank(value):
        raise ValueError(f"Missing numeric value for {label}.")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"Non-finite numeric value for {label}: {value}")
    return result


def as_int(value: Any, label: str) -> int:
    return int(round(as_float(value, label)))


def require_columns(df: pd.DataFrame, required: list[str], source: str) -> None:
    missing = [column for column in required if column not in df.columns]
    if missing:
        raise KeyError(
            f"Missing columns in {source}: {missing}. "
            f"Available columns: {list(df.columns)}"
        )


def covariate_label(value: Any) -> str:
    if is_blank(value) or str(value).strip() == "(none)":
        return "None"

    covariates = [part.strip() for part in str(value).split(",") if part.strip()]
    return "; ".join(COVARIATE_LABELS.get(item, item) for item in covariates)


def pair_label(pair: str, data: pd.DataFrame) -> str:
    subset = data.loc[data["Pair"].astype(str).str.strip() == pair]
    if subset.empty:
        raise ValueError(f"Pair {pair!r} was not found in data3.xlsx.")

    targets = subset["V_Target"].astype(str).str.strip().dropna().unique()
    leaders = subset["V_Leading_Class"].astype(str).str.strip().dropna().unique()

    if len(targets) != 1 or len(leaders) != 1:
        raise ValueError(
            f"Pair {pair!r} maps to multiple target/leader classes: "
            f"targets={targets.tolist()}, leaders={leaders.tolist()}"
        )

    return f"{targets[0]} → {leaders[0]}"


def find_unique(
    df: pd.DataFrame,
    pair: str,
    distribution: str,
    source: str,
) -> pd.Series:
    mask = (
        df["Pair"].astype(str).str.strip().eq(pair)
        & df["Distribution"].astype(str).str.strip().eq(distribution)
    )
    matches = df.loc[mask]

    if len(matches) != 1:
        raise ValueError(
            f"Expected exactly one row for pair={pair!r}, "
            f"distribution={distribution!r} in {source}; found {len(matches)}."
        )

    return matches.iloc[0]


def stat_string(stat: Any, p_value: Any, bootstrap_status: Any) -> str:
    status = "" if is_blank(bootstrap_status) else str(bootstrap_status).strip()

    if is_blank(p_value):
        if status and status != "OK":
            return "Bootstrap invalid"
        return "Not available"

    return f"{as_float(stat, 'test statistic'):.3f} ({as_float(p_value, 'bootstrap p'):.3f})"


def decision_label(decision_row: pd.Series) -> str:
    selected = decision_row.get("Selected distribution")
    if is_blank(selected):
        return "Not adequate"

    selected_name = str(selected).strip()
    lowest_name = str(decision_row.get("Lowest-AIC tested distribution", "")).strip()

    if selected_name != lowest_name:
        return "Accepted alternative"
    return "Accepted"


# =============================================================================
# BUILD TABLE 8
# =============================================================================
def build_table8(
    t5_path: Path,
    t6_path: Path,
    data_path: Path,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    for path in (t5_path, t6_path, data_path):
        if not path.exists():
            raise FileNotFoundError(f"Input file not found: {path}")

    t5_models = pd.read_excel(t5_path, sheet_name="S1_Model_summary")
    t6_gof = pd.read_excel(t6_path, sheet_name="S3_AFT_final_GOF")
    t6_decisions = pd.read_excel(t6_path, sheet_name="S5_Final_decisions")
    data = pd.read_excel(data_path, sheet_name="Sheet1")

    require_columns(
        t5_models,
        [
            "Pair",
            "Distribution",
            "n",
            "Baseline AIC",
            "Final AIC",
            "Total ΔAIC",
            "Final covariates",
        ],
        "T5/S1_Model_summary",
    )
    require_columns(
        t6_gof,
        [
            "Pair",
            "Distribution",
            "Final covariates",
            "n",
            "AIC",
            "ΔAIC",
            "KS D",
            "KS bootstrap p",
            "AD A²",
            "AD bootstrap p",
            "Bootstrap status",
        ],
        "T6/S3_AFT_final_GOF",
    )
    require_columns(
        t6_decisions,
        [
            "Pair",
            "Selected distribution",
            "Selected covariates",
            "Lowest-AIC tested distribution",
            "Lowest-AIC test outcome",
        ],
        "T6/S5_Final_decisions",
    )
    require_columns(data, ["V_Target", "V_Leading_Class", "Pair"], "data3/Sheet1")

    decisions_by_pair = {
        str(row["Pair"]).strip(): row
        for _, row in t6_decisions.iterrows()
    }

    table_rows: list[dict[str, Any]] = []
    audit_rows: list[dict[str, Any]] = []

    for pair in PAIR_ORDER:
        if pair not in decisions_by_pair:
            raise KeyError(f"Pair {pair!r} is absent from T6/S5_Final_decisions.")

        decision = decisions_by_pair[pair]
        selected_distribution = decision.get("Selected distribution")
        lowest_distribution = decision.get("Lowest-AIC tested distribution")

        displayed_distribution = (
            str(selected_distribution).strip()
            if not is_blank(selected_distribution)
            else str(lowest_distribution).strip()
        )

        t6_row = find_unique(
            t6_gof,
            pair=pair,
            distribution=displayed_distribution,
            source="T6/S3_AFT_final_GOF",
        )
        t5_row = find_unique(
            t5_models,
            pair=pair,
            distribution=displayed_distribution,
            source="T5/S1_Model_summary",
        )

        n_data = int((data["Pair"].astype(str).str.strip() == pair).sum())
        n_t5 = as_int(t5_row["n"], f"T5 n for {pair}")
        n_t6 = as_int(t6_row["n"], f"T6 n for {pair}")

        validation = "OK" if n_data == n_t5 == n_t6 else "N MISMATCH"
        if validation != "OK":
            raise ValueError(
                f"Sample-size mismatch for {pair}: "
                f"data3={n_data}, T5={n_t5}, T6={n_t6}."
            )

        baseline_aic = as_float(t5_row["Baseline AIC"], f"T5 baseline AIC for {pair}")
        final_aic_t5 = as_float(t5_row["Final AIC"], f"T5 final AIC for {pair}")
        stored_improvement = as_float(t5_row["Total ΔAIC"], f"T5 Total ΔAIC for {pair}")
        recalculated_improvement = baseline_aic - final_aic_t5

        if not math.isclose(
            stored_improvement,
            recalculated_improvement,
            rel_tol=0,
            abs_tol=1e-8,
        ):
            raise ValueError(
                f"AIC-improvement inconsistency for {pair}, {displayed_distribution}: "
                f"stored={stored_improvement}, recalculated={recalculated_improvement}."
            )

        bootstrap_status = t6_row.get("Bootstrap status")
        retained_covariates = t6_row.get("Final covariates")

        table_rows.append(
            {
                "Target → leader": pair_label(pair, data),
                "Selected distribution": displayed_distribution,
                "Retained covariates": covariate_label(retained_covariates),
                "AIC improvement over covariate-free model": round(stored_improvement, 3),
                "Final AIC": round(as_float(t6_row["AIC"], f"T6 AIC for {pair}"), 3),
                "KS D (p)": stat_string(
                    t6_row["KS D"],
                    t6_row["KS bootstrap p"],
                    bootstrap_status,
                ),
                "AD A² (p)": stat_string(
                    t6_row["AD A²"],
                    t6_row["AD bootstrap p"],
                    bootstrap_status,
                ),
                "Decision": decision_label(decision),
            }
        )

        audit_rows.append(
            {
                "Pair key": pair,
                "Target → leader": pair_label(pair, data),
                "n from data3": n_data,
                "n in T5": n_t5,
                "n in T6": n_t6,
                "Displayed distribution": displayed_distribution,
                "T6 selected distribution": (
                    None if is_blank(selected_distribution) else str(selected_distribution).strip()
                ),
                "Lowest-AIC tested distribution": str(lowest_distribution).strip(),
                "T5 baseline AIC": baseline_aic,
                "T5 final AIC": final_aic_t5,
                "T5 total ΔAIC": stored_improvement,
                "T6 final AIC": as_float(t6_row["AIC"], f"T6 AIC for {pair}"),
                "T6 ΔAIC within final evidence set": as_float(
                    t6_row["ΔAIC"], f"T6 ΔAIC for {pair}"
                ),
                "Bootstrap status": bootstrap_status,
                "Validation": validation,
            }
        )

    return pd.DataFrame(table_rows), pd.DataFrame(audit_rows)


# =============================================================================
# EXPORT WITH OPENPYXL FORMATTING
# =============================================================================
def format_workbook(output_path: Path) -> None:
    wb = load_workbook(output_path)
    ws = wb["Table 8"]
    audit = wb["Audit"]

    # Insert title and spacing above the dataframe table.
    ws.insert_rows(1, amount=2)
    ws.merge_cells("A1:H1")
    ws["A1"] = "Table 8. Final Covariate-Conditioned AFT Models by Target–Leader Pair"

    title_fill = PatternFill("solid", fgColor="1F4E78")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    note_fill = PatternFill("solid", fgColor="F7F7F7")
    thin = Side(style="thin", color="A6A6A6")
    hair = Side(style="hair", color="B7B7B7")

    ws["A1"].fill = title_fill
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for cell in ws[3]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[3].height = 36

    for row in ws.iter_rows(min_row=4, max_row=11, min_col=1, max_col=8):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=hair, right=hair, top=hair, bottom=hair)
    for row_number in range(4, 12):
        ws.row_dimensions[row_number].height = 30

    # Number formats.
    for row_number in range(4, 12):
        ws.cell(row=row_number, column=4).number_format = "0.000"
        ws.cell(row=row_number, column=5).number_format = "0.000"

    # Add note below the table.
    note_row = 13
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=8)
    ws.cell(note_row, 1).value = (
        "Note: AIC improvement is calculated against the covariate-free model of the same "
        "distribution using T5. Final AIC and bootstrap goodness-of-fit results are taken from "
        "T6. Sample size is independently verified against data3. Acceptance requires "
        "bootstrap KS p ≥ 0.05 and AD p ≥ 0.05. For unresolved pairs, the lowest-AIC tested "
        "distribution is displayed and marked Not adequate."
    )
    ws.cell(note_row, 1).font = Font(italic=True, size=9)
    ws.cell(note_row, 1).fill = note_fill
    ws.cell(note_row, 1).alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 18,
        "B": 22,
        "C": 38,
        "D": 20,
        "E": 12,
        "F": 16,
        "G": 16,
        "H": 21,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A4"

    # Decision highlighting.
    green_fill = PatternFill("solid", fgColor="E2F0D9")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    red_fill = PatternFill("solid", fgColor="FCE4D6")

    ws.conditional_formatting.add(
        "H4:H11",
        FormulaRule(formula=['H4="Accepted"'], fill=green_fill),
    )
    ws.conditional_formatting.add(
        "H4:H11",
        FormulaRule(formula=['H4="Accepted alternative"'], fill=yellow_fill),
    )
    ws.conditional_formatting.add(
        "H4:H11",
        FormulaRule(formula=['H4="Not adequate"'], fill=red_fill),
    )

    # Audit formatting.
    for cell in audit[1]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in audit.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for column_index in range(1, audit.max_column + 1):
        letter = get_column_letter(column_index)
        max_length = 0
        for cell in audit[letter]:
            text = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(text))
        audit.column_dimensions[letter].width = min(max(max_length + 2, 10), 28)

    audit.freeze_panes = "A2"

    wb.save(output_path)


def export_table8(table8: pd.DataFrame, audit: pd.DataFrame) -> None:
    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        table8.to_excel(writer, sheet_name="Table 8", index=False)
        audit.to_excel(writer, sheet_name="Audit", index=False)

    format_workbook(OUTPUT_XLSX)
    table8.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")

    print("Generated:")
    print(OUTPUT_XLSX)
    print(OUTPUT_CSV)


if __name__ == "__main__":
    table8_df, audit_df = build_table8(T5_PATH, T6_PATH, DATA_PATH)
    export_table8(table8_df, audit_df)
